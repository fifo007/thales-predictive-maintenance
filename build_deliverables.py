"""Create the two requested Word deliverables from the supplied dataset."""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).parent
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
df = pd.read_csv(ROOT / "Thales_Group_Manufacturing.csv")
df["event_time"] = pd.to_datetime(df.Date + " " + df.Timestamp, dayfirst=True)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)
def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "D9D9D9"); b.append(e)
    tcPr.append(b)
def style(doc):
    section = doc.sections[0]; section.top_margin=Inches(.72); section.bottom_margin=Inches(.72); section.left_margin=Inches(.8); section.right_margin=Inches(.8)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos"); normal.font.size=Pt(10)
    for name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 12)]:
        s=doc.styles[name]; s.font.name="Aptos Display"; s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display"); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
def add_title(doc, title, subtitle):
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.add_run(title)
    p=doc.add_paragraph(); p.add_run(subtitle).italic=True; p.runs[0].font.color.rgb=RGBColor(89,89,89)
def add_table(doc, headers, rows):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style="Table Grid"
    for i,h in enumerate(headers):
        c=table.rows[0].cells[i]; c.text=str(h); shade(c,"17365D"); border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    for j,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); border(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j % 2 == 1: shade(cells[i],"F3F7FB")
            for r in cells[i].paragraphs[0].runs: r.font.size=Pt(9)
    doc.add_paragraph()
    return table
def bullets(doc, values):
    for v in values: doc.add_paragraph(v, style="List Bullet")

start, end = df.event_time.min(), df.event_time.max()
machines = df.Machine_ID.nunique(); rows=len(df)
status = df.Efficiency_Status.value_counts(); modes = df.Operation_Mode.value_counts()

paper=Document(); style(paper)
add_title(paper, "Predictive Maintenance Analytics Research Paper", "Thales Manufacturing Telemetry Study")
paper.add_paragraph(f"This study analyzes {rows:,} telemetry readings across {machines} industrial machines, covering {start:%d %b %Y %H:%M} to {end:%d %b %Y %H:%M}. The recommended approach is machine-specific baseline modelling with multivariate anomaly detection so maintenance teams can identify early warning signals before costly breakdowns.")
paper.add_heading("1 Background and problem", level=1)
paper.add_paragraph("Manufacturing failures are relatively rare but expensive. Waiting for broad efficiency measures to degrade can be too late because temperature, vibration, power, and error behavior may shift subtly beforehand. Fixed thresholds are insufficient because normal operating behavior differs by machine and operation mode.")
paper.add_heading("2 Exploratory data analysis", level=1)
add_table(paper, ["Data characteristic", "Finding"], [
    ["Observations", f"{rows:,} timestamped readings"], ["Assets", f"{machines} unique machine identifiers"],
    ["Operation modes", ", ".join(f"{k}: {v:,}" for k,v in modes.items())],
    ["Efficiency labels", ", ".join(f"{k}: {v:,}" for k,v in status.items())],
    ["Core sensor domains", "Temperature, vibration, power, connectivity, quality, production speed, maintenance score, and error rate"],
])
paper.add_heading("3 Methodology", level=1)
paper.add_heading("Baseline behavior modelling", level=2)
paper.add_paragraph("For each Machine_ID, the application computes rolling median baselines for temperature, vibration, power consumption, and error rate. Median baselines are robust to occasional spikes. Each reading is represented by its normalized absolute deviation from the recent machine baseline.")
paper.add_heading("Feature engineering", level=2)
bullets(paper, ["Rolling sensor deviations for temperature, vibration, power, and error rate.", "Vibration-to-power ratio to capture mechanical instability under changing load.", "Ten-observation error-rate trend to capture deterioration.", "Original sensor, network, quality, production, and maintenance-score fields for multivariate context."])
paper.add_heading("Anomaly detection and risk scoring", level=2)
paper.add_paragraph("Isolation Forest is trained on a deterministic representative sample of the telemetry. It isolates unusual combinations of measurements without requiring a confirmed failure label. Raw anomaly values are normalized to a 0-100 score. Scores up to 45 are Low risk, 45-70 Medium risk, and above 70 High risk; the dashboard allows users to change the operational alert threshold.")
paper.add_heading("4 Findings and operational recommendations", level=1)
bullets(paper, ["Use the dashboard's high-risk asset list as the daily inspection queue, not as an autonomous shutdown instruction.", "Investigate sustained upward anomaly trends before they become high-risk alerts; this is the available early-warning lead time.", "Compare sensor patterns during Maintenance mode against Active and Idle behavior to validate stabilization after intervention.", "Review threshold performance with maintenance outcomes monthly, and recalibrate contamination and risk cutoffs as confirmed-failure labels become available."])
paper.add_heading("5 Limitations and next steps", level=1)
paper.add_paragraph("The supplied data does not include confirmed breakdown dates, work orders, costs, or inspection outcomes. The model therefore identifies behavioral anomalies rather than a calibrated probability of failure. The next deployment phase should join telemetry to maintenance records, measure precision and lead time, and establish a human-approved response playbook.")
paper.add_heading("6 Conclusion", level=1)
paper.add_paragraph("Machine-specific baselines and multivariate anomaly scores provide a practical bridge from reactive maintenance to preventive inspection. The accompanying Streamlit application turns this method into a live analytical workflow with transparent filters, alert export, and historical risk review.")
paper.save(OUT / "predictive_maintenance_research_paper.docx")

summary=Document(); style(summary)
add_title(summary, "Executive Summary for Government Stakeholders", "Predictive Maintenance Analytics for Industrial Resilience")
summary.add_paragraph("The proposed dashboard uses operational telemetry from industrial machines to identify unusual behavior before it becomes a breakdown. It supports more reliable production, more targeted maintenance activity, and evidence-based oversight of digital industrial infrastructure.")
summary.add_heading("What has been delivered", level=1)
add_table(summary,["Deliverable", "Value"],[
    ["Live Streamlit dashboard", "Risk distribution, machine-level anomaly trends, maintenance alerts, and historical escalation analysis."],
    ["Risk classification", "Low, Medium, and High categories that turn complex sensor patterns into inspection priorities."],
    ["User controls", "Machine selection, risk-threshold adjustment, time-window selection, and operation-mode filtering."],
    ["Exportable evidence", "Filtered alert list can be downloaded for maintenance coordination and audit trails."],
])
summary.add_heading("Why it matters", level=1)
bullets(summary,["Reduces reliance on after-the-fact maintenance by surfacing early sensor deviations.", "Supports continuity of essential manufacturing operations by prioritizing assets that need attention.", "Creates a transparent analytical record that can be reviewed alongside maintenance decisions.", "Strengthens digital resilience by considering connectivity and packet-loss signals with physical sensors."])
summary.add_heading("Governance recommendations", level=1)
bullets(summary,["Keep a qualified maintenance professional in the decision loop; the dashboard prioritizes inspection and does not initiate shutdowns.", "Maintain data-quality checks for missing, delayed, or implausible sensor values.", "Track inspection outcomes, failures avoided, downtime, and cost to validate benefits and improve thresholds.", "Apply role-based access and retention policies before production deployment, particularly where telemetry is operationally sensitive."])
summary.add_heading("Decision requested", level=1)
summary.add_paragraph("Authorize a controlled pilot that links dashboard alerts to maintenance work orders. Success should be measured through warning lead time, percentage of high-risk alerts inspected, avoided downtime, and false-alert rate. The pilot should report results before wider deployment.")
summary.save(OUT / "executive_summary_government_stakeholders.docx")
print("Created deliverables")
